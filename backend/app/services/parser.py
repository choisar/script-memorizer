import io
import re
from typing import List, Tuple
import pypdfium2 as pdfium
import pdfplumber
import docx
from app.models.schemas import ScriptSegment

# 문장 분할 정규식
# 마침표/물음표/느낌표 뒤에 공백이나 줄바꿈이 올 때 분할하되,
# 번호 매기기(1., 2., 10.), 소수점(3.14), 말줄임표(..) 등에서 오분할 방지
SENTENCE_SPLIT_REGEX = re.compile(
    r'(?<=[.?!])(?<!\b\d\.)(?<!\b\d\d\.)(?<!\b[0-9]\.[0-9])(?<!\.\.)\s+(?=[A-Za-z0-9가-힣"\'“‘])'
)

ITEM_HEADER_REGEX = re.compile(
    r'^(\d+\.\s+|\d+\)\s*|\(\d+\)\s*|[가-하]\.\s*|[①-⑳]\s*|#{1,6}\s+|\[.*?\])'
)

ITEM_BODY_OR_NUM_REGEX = re.compile(
    r'^(\d+\.\s+|\d+\)\s*|\(\d+\)\s*|\d+\s+[가-힣A-Za-z]|[가-하]\.\s*|[①-⑳]\s*|#{1,6}\s+|\[.*?\]|[가-힣A-Za-z0-9_]{1,8}:)'
)

def unwrap_pdf_lines(text: str) -> str:
    """PDF에서 페이지 폭에 의해 인위적으로 강제 줄바꿈된 문장을 자연스럽게 합쳐줍니다."""
    raw_lines = [l.strip() for l in text.replace('\r\n', '\n').split('\n') if l.strip()]
    merged_lines = []

    for line in raw_lines:
        if not merged_lines:
            merged_lines.append(line)
            continue

        prev = merged_lines[-1]
        prev_is_header = bool(ITEM_HEADER_REGEX.match(prev))
        prev_ends_punct = bool(re.search(r'[.?!:]$', prev))
        curr_is_item = bool(ITEM_BODY_OR_NUM_REGEX.match(line))

        # 앞줄이 제목/항목 헤더가 아니고, 문장부호로 끝나지 않았으며, 뒷줄이 새로운 항목/번호가 아닌 경우 개행 병합
        if not prev_is_header and not prev_ends_punct and not curr_is_item:
            merged_lines[-1] = prev + ' ' + line
        else:
            merged_lines.append(line)

    return "\n".join(merged_lines)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    pypdfium2(Chromium 기반 PDFium)를 이용해 텍스트 레이아웃 순서대로 정교하게 추출하고,
    인위적인 줄바꿈을 복원합니다.
    """
    try:
        pdf = pdfium.PdfDocument(file_bytes)
        page_texts = []
        for page in pdf:
            textpage = page.get_textpage()
            page_text = textpage.get_text_range()
            if page_text and page_text.strip():
                unwrapped = unwrap_pdf_lines(page_text)
                if unwrapped:
                    page_texts.append(unwrapped)
        return "\n\n".join(page_texts)
    except Exception:
        # Fallback to pdfplumber if pdfium fails
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
        return "\n\n".join(text_parts)

def extract_text_from_docx(file_bytes: bytes) -> str:
    """DOCX 바이트로부터 텍스트를 추출합니다."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)

def extract_text_from_txt(file_bytes: bytes) -> str:
    """TXT/MD 바이트로부터 텍스트를 디코딩합니다 (UTF-8, CP949/EUC-KR 자동 감지)."""
    for encoding in ["utf-8", "utf-8-sig", "cp949", "euc-kr"]:
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="ignore")

def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """확장자에 따라 적절한 텍스트 추출 함수를 호출합니다."""
    ext = filename.lower().split(".")[-1]
    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in ["docx", "doc"]:
        return extract_text_from_docx(file_bytes)
    elif ext in ["txt", "md", "markdown"]:
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {ext} (pdf, docx, txt, md 지원)")

def split_into_sentences(paragraph_text: str) -> List[str]:
    """문단을 문장 단위로 분할합니다. 줄바꿈 및 문장 부호(.?!)를 정교하게 분리합니다."""
    clean_text = paragraph_text.strip()
    if not clean_text:
        return []

    lines = [line.strip() for line in clean_text.split('\n') if line.strip()]
    result: List[str] = []

    for line in lines:
        parts = SENTENCE_SPLIT_REGEX.split(line)
        for part in parts:
            p_clean = part.strip()
            if p_clean:
                result.append(p_clean)

    if not result:
        result = [clean_text]

    return result

def clean_markdown_line(line: str) -> str:
    """마크다운 헤더 기호(##, ### 등) 및 볼드(**) 기호를 정리하여 발화하기 좋은 텍스트로 다듬습니다."""
    # 선두 헤딩 기호 제거
    line = re.sub(r'^#{1,6}\s*', '', line).strip()
    # 마크다운 볼드/이탤릭 제거 (**굵게** -> 굵게)
    line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
    line = re.sub(r'\*(.*?)\*', r'\1', line)
    return line.strip()

def split_into_segments(full_text: str) -> List[ScriptSegment]:
    """
    전체 텍스트를 [문단 > 문장] 계층 구조의 ScriptSegment 리스트로 변환합니다.
    - 마크다운 대/중제목(#, ##)이 존재하는 경우, 대/중제목 단위로 문단을 묶어 하위 소제목(###) 및 항목들이 동일 문단 내 세부 문장으로 귀속되도록 합니다.
    - 대/중제목은 없지만 1. , 2. 와 같은 주요 번호 매기기가 2개 이상 존재하는 경우 번호 단위로 상위 문단을 분할합니다.
    - 그 외의 경우, 줄바꿈 2회 이상(\n\n)을 문단 구분으로 처리합니다.
    - 각 문단 내에서는 줄바꿈 및 문장 부호(.?!)를 기준으로 세부 문장으로 분할합니다.
    """
    normalized_text = full_text.replace("\r\n", "\n").replace("\r", "\n")

    # 1. 문서에 마크다운 대/중제목(#, ##)이 포함되어 있으면 대/중제목 단위로 문단 분할
    if re.search(r'(?m)^#{1,2}\s+', normalized_text):
        raw_chunks = re.split(r'(?m)(?=^#{1,2}\s+)', normalized_text)
        raw_paragraphs = [c.strip() for c in raw_chunks if c.strip()]
    # 2. 대/중제목은 없지만 1. , 2. 와 같은 주요 번호 매기기가 복수 개 존재하는 경우 번호 단위로 문단 분할
    elif len(re.findall(r'(?m)^\d+\.\s+', normalized_text)) >= 2:
        raw_chunks = re.split(r'(?m)(?=^\d+\.\s+)', normalized_text)
        raw_paragraphs = [c.strip() for c in raw_chunks if c.strip()]
    else:
        # 3. 일반 대본/텍스트의 경우 빈 줄(\n\n) 기준으로 문단 분할
        raw_paragraphs = [p.strip() for p in re.split(r'\n\s*\n', normalized_text) if p.strip()]



    segments: List[ScriptSegment] = []

    for p_idx, p_raw in enumerate(raw_paragraphs, start=1):
        clean_lines = [clean_markdown_line(l) for l in p_raw.split('\n') if l.strip()]
        p_text = "\n".join(clean_lines).strip()
        if not p_text:
            continue

        p_id = f"p{p_idx}"
        sentence_texts = split_into_sentences(p_text)

        sentence_segments: List[ScriptSegment] = []
        for s_idx, s_text in enumerate(sentence_texts, start=1):
            s_id = f"{p_id}-s{s_idx}"
            sentence_segments.append(
                ScriptSegment(
                    id=s_id,
                    type="sentence",
                    content=s_text,
                    children=[]
                )
            )

        segments.append(
            ScriptSegment(
                id=p_id,
                type="paragraph",
                content=p_text,
                children=sentence_segments
            )
        )

    return segments

